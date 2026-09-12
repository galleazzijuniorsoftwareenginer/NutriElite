"""Builds the downloadable spreadsheet for a plan's weekly micronutrient
breakdown (calculate_plan_micronutrients' output) — the "planilla de
micronutrientes" the nutritionist can save/print/share, mirroring the PDF
export already offered for the clinical report (see pdf_service.py)."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="EFEBFF", end_color="EFEBFF", fill_type="solid")
HEADER_FONT = Font(bold=True, color="4A37D1")
TOTAL_FONT = Font(bold=True)


def generate_micronutrients_xlsx(result: dict) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Micronutrientes"

    dias = result["dias"]
    campos = result["campos"]
    fields = list(campos.keys())

    headers = ["Nutriente", "Unidad"] + [d["dia"] for d in dias] + ["Total semana"]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for row_i, field in enumerate(fields, start=2):
        ws.cell(row=row_i, column=1, value=campos[field]["label"])
        ws.cell(row=row_i, column=2, value=campos[field]["unidad"])
        for col_i, day in enumerate(dias, start=3):
            value = day["totales"].get(field)
            ws.cell(row=row_i, column=col_i, value=round(value, 2) if value is not None else None)
        total_cell = ws.cell(row=row_i, column=len(dias) + 3, value=round(result["totales_semana"][field], 2))
        total_cell.font = TOTAL_FONT

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.column_dimensions["A"].width = 22

    # Hoja de cobertura: qué tan completos son estos datos, sin esconderlo.
    ws2 = wb.create_sheet("Cobertura de datos")
    ws2.cell(row=1, column=1, value="Ingredientes con dato USDA").font = HEADER_FONT
    ws2.cell(row=1, column=2, value=f"{result['cobertura']['con_datos']}/{result['cobertura']['total']}")
    ws2.cell(row=3, column=1, value="Ingredientes sin dato disponible en USDA:").font = Font(bold=True)
    for i, name in enumerate(result["ingredientes_sin_datos"], start=4):
        ws2.cell(row=i, column=1, value=f"• {name}")
    ws2.column_dimensions["A"].width = 45

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
